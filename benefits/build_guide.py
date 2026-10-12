"""Build the candidate benefits guide. Light letterhead, logo and contact on every page."""

from pathlib import Path
import json

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
INK = RGBColor(0x1C, 0x24, 0x30)
GOLD = RGBColor(0x8A, 0x73, 0x54)
MUTED = RGBColor(0x5C, 0x65, 0x70)
PAPER = "FBF8F4"
STONE = "F4EFE6"
RULE = "E4D9C8"
CONTENT = 7.1


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    props.append(shading)


def cell_margins(cell, dxa: str = "80") -> None:
    props = cell._tc.get_or_add_tcPr()
    inset = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), dxa)
        node.set(qn("w:type"), "dxa")
        inset.append(node)
    props.append(inset)


def cell_borders(cell, *, color=RULE, size="4", edges=("top", "left", "bottom", "right")) -> None:
    props = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        if edge in edges:
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), color)
        else:
            node.set(qn("w:val"), "nil")
        borders.append(node)
    props.append(borders)


def valign(cell, value: str = "center") -> None:
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:vAlign")
    node.set(qn("w:val"), value)
    props.append(node)


def set_width(cell, inches: float) -> None:
    props = cell._tc.get_or_add_tcPr()
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")
    props.append(width)


def table_width(table, inches: float) -> None:
    props = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")
    props.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(layout)


def tighten(paragraph, *, before=0, after=0, line=1.08) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run(run, *, size=11, bold=False, color=INK, font="Calibri") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_text(paragraph, text, **style) -> None:
    set_run(paragraph.add_run(text), **style)


def page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def paint_page(document) -> None:
    background = OxmlElement("w:background")
    background.set(qn("w:color"), PAPER)
    document.element.insert(0, background)
    shown = OxmlElement("w:displayBackgroundShape")
    document.settings.element.append(shown)


def write_header(section, guide: dict) -> None:
    header = section.header
    header.is_linked_to_previous = False
    opening = header.paragraphs[0]
    opening.text = ""
    tighten(opening)
    grid = header.add_table(rows=1, cols=3, width=Inches(CONTENT))
    table_width(grid, CONTENT)
    logo, identity, contact = grid.rows[0].cells
    widths = (0.58, 3.55, 2.97)
    for cell, width in zip(grid.rows[0].cells, widths):
        set_width(cell, width)
        shade(cell, PAPER)
        cell_borders(cell, edges=())
        cell_margins(cell, "40")
        valign(cell, "center")
    logo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    tighten(logo.paragraphs[0])
    logo.paragraphs[0].add_run().add_picture(str(ROOT / "images" / "mark.jpg"), width=Inches(0.42))
    identity.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    tighten(identity.paragraphs[0], after=1)
    add_text(identity.paragraphs[0], guide["company"], size=12, bold=True)
    season = identity.add_paragraph()
    tighten(season)
    add_text(season, f"Benefits  {guide['season']}", size=9, color=GOLD)
    contact.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tighten(contact.paragraphs[0], after=0)
    add_text(contact.paragraphs[0], guide["address"], size=8, color=MUTED)
    phones = contact.add_paragraph()
    phones.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tighten(phones, after=0)
    add_text(phones, guide["phone_office"], size=8, color=INK)
    mail = contact.add_paragraph()
    mail.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tighten(mail)
    add_text(mail, guide["email"], size=8, color=GOLD)
    rule = header.add_paragraph()
    tighten(rule, before=4)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "8A7354")
    border.append(bottom)
    rule._p.get_or_add_pPr().append(border)


def keep_row(row) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))


def write_footer(section, guide: dict) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    opening = footer.paragraphs[0]
    opening.text = ""
    tighten(opening)
    grid = footer.add_table(rows=1, cols=2, width=Inches(CONTENT))
    table_width(grid, CONTENT)
    left, right = grid.rows[0].cells
    set_width(left, 5.2)
    set_width(right, 1.9)
    for cell in (left, right):
        shade(cell, PAPER)
        cell_borders(cell, color=RULE, size="8", edges=("top",))
        cell_margins(cell, "40")
    tighten(left.paragraphs[0], before=2)
    add_text(left.paragraphs[0], f"For candidates   ·   {guide['season']}", size=8, color=MUTED)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tighten(right.paragraphs[0], before=2)
    add_text(right.paragraphs[0], "Page ", size=8, color=MUTED)
    page_field(right.paragraphs[0])


def plain_text(guide: dict) -> str:
    lines = [
        f"{guide['company']} Employee Benefits Summary",
        f"Plan year {guide['season']}",
        "",
        guide["address"],
        f"Phone: {guide['phone_office']} (office). Email: {guide['email']}",
        "",
        guide["intro"],
        "",
    ]
    for item in guide["glance"]:
        lines.append(f"{item['value']} — {item['label']}")
    for section in guide["sections"]:
        lines.extend(["", section.get("kicker", ""), section["title"], section["lead"]])
        lines.extend(f"• {point}" for point in section["points"])
        table = section.get("table")
        if table:
            lines.append(" | ".join(table["columns"]))
            lines.extend(" | ".join(row) for row in table["rows"])
    lines.extend(["", f"Source: {guide['source']}"])
    return "\n".join(lines) + "\n"


def build() -> None:
    guide = json.loads((ROOT / "benefits.json").read_text())
    (ROOT / "Janus-Soft-Benefits-2026-27.txt").write_text(plain_text(guide))
    document = Document()
    paint_page(document)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.45)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.28)
    section.different_first_page_header_footer = False
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    write_header(section, guide)
    write_footer(section, guide)

    eyebrow = document.add_paragraph()
    tighten(eyebrow, after=2)
    add_text(eyebrow, "FOR CANDIDATES", size=9, bold=True, color=GOLD)
    title = document.add_paragraph()
    tighten(title, after=2)
    add_text(title, "Your benefits", size=28, bold=True)
    intro = document.add_paragraph()
    tighten(intro, before=6, after=10, line=1.15)
    add_text(intro, guide["intro"], size=11, color=MUTED)

    glance = document.add_table(rows=1, cols=4)
    table_width(glance, CONTENT)
    for index, item in enumerate(guide["glance"]):
        cell = glance.cell(0, index)
        set_width(cell, CONTENT / 4)
        shade(cell, STONE)
        cell_margins(cell, "90")
        cell_borders(cell, color="8A7354", size="8", edges=("top",))
        cell_borders_rest = cell
        # Top rule is gold; the other edges stay quiet.
        props = cell._tc.get_or_add_tcPr()
        borders = props.find(qn("w:tcBorders"))
        for edge, color, size in (("left", RULE, "4"), ("right", RULE, "4"), ("bottom", RULE, "4")):
            node = borders.find(qn(f"w:{edge}"))
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), color)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        tighten(cell.paragraphs[0], after=2)
        add_text(cell.paragraphs[0], item["value"], size=18, bold=True, color=GOLD)
        label = cell.add_paragraph()
        tighten(label)
        add_text(label, item["label"], size=8, color=MUTED)
        del cell_borders_rest

    note = document.add_paragraph()
    tighten(note, before=8, after=4)
    add_text(note, "The section that follows states the term. This guide does not add coverage beyond the summary.", size=9, color=MUTED)

    for block in guide["sections"]:
        if block.get("break_before"):
            document.add_page_break()
        else:
            gap = document.add_paragraph()
            tighten(gap, before=2, after=0)
        row = document.add_table(rows=1, cols=2)
        table_width(row, CONTENT)
        icon, body = row.cell(0, 0), row.cell(0, 1)
        set_width(icon, 1.15)
        set_width(body, CONTENT - 1.15)
        for cell in (icon, body):
            shade(cell, PAPER)
            cell_borders(cell, edges=())
            cell_margins(cell, "40")
            valign(cell, "top")
        icon.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tighten(icon.paragraphs[0])
        icon.paragraphs[0].add_run().add_picture(str(ROOT / "images" / block["image"]), width=Inches(0.8))
        tighten(body.paragraphs[0], after=0)
        add_text(body.paragraphs[0], block["kicker"], size=9, bold=True, color=GOLD)
        heading = body.add_paragraph()
        tighten(heading, before=1, after=2)
        heading.paragraph_format.keep_with_next = True
        add_text(heading, block["title"], size=18, bold=True)
        lead = body.add_paragraph()
        tighten(lead, before=1, after=4, line=1.08)
        add_text(lead, block["lead"], size=11, color=MUTED)
        for point in block["points"]:
            bullet = body.add_paragraph()
            tighten(bullet, before=1, after=2, line=1.08)
            add_text(bullet, "–  ", size=11, color=GOLD)
            add_text(bullet, point, size=11)

        table = block.get("table")
        if not table:
            continue
        grid = document.add_table(rows=1 + len(table["rows"]), cols=len(table["columns"]))
        table_width(grid, CONTENT)
        shares = (1.45, 1.55, 1.55, 2.55)
        for index, name in enumerate(table["columns"]):
            header = grid.rows[0].cells[index]
            set_width(header, shares[index])
            shade(header, STONE)
            cell_margins(header, "70")
            cell_borders(header, color="8A7354", size="8", edges=("bottom",))
            tighten(header.paragraphs[0])
            header.paragraphs[0].paragraph_format.keep_with_next = True
            add_text(header.paragraphs[0], name, size=8, bold=True, color=INK)
        keep_row(grid.rows[0])
        for r_index, values in enumerate(table["rows"]):
            for c_index, value in enumerate(values):
                data = grid.rows[r_index + 1].cells[c_index]
                set_width(data, shares[c_index])
                shade(data, "FFFFFF" if r_index % 2 == 0 else STONE)
                cell_margins(data, "70")
                cell_borders(data, color=RULE, size="4", edges=("bottom",))
                tighten(data.paragraphs[0])
                add_text(data.paragraphs[0], value, size=10, bold=c_index == 0)
            keep_row(grid.rows[r_index + 1])

    close = document.add_paragraph()
    tighten(close, before=16)
    add_text(close, f"These terms are the {guide['season']} benefits summary of {guide['company']}.", size=8, color=MUTED)
    document.core_properties.title = f"{guide['company']} Benefits {guide['season']}"
    document.core_properties.subject = "Benefits guide for candidates"
    document.core_properties.category = "Benefits"
    document.save(ROOT / guide["filename"])


if __name__ == "__main__":
    build()
