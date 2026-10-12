"""Section-aware knowledge answers, the written-answer check, and the résumé check before applying."""

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from app.core.extract import extract_text
from app.knowledge.graph import blocks, figures_match, quoted_answer, readable
from app.knowledge.intent import job_part, names_job, wants_knowledge
from app.knowledge.sections import chunk_sections, clean_lines
from app.main import create_app
from app.models import Job
from app.public.resume_check import reset_check_limit
from app.public.routes import reset_rate_limit
from tests.fakes import RecordingEmbedder, RecordingLLM

HEADER = "Janus Soft Inc.\nBenefits 2026–27\n2700 Prosperity Ave, Fairfax VA"
GUIDE = f"""{HEADER}
FOR CANDIDATES
Your benefits
Prepared for candidates considering Janus Soft. Every figure is the term in the company summary.
100%
Employee health premium
Unlimited
Paid time off

01 · Care
Medical, dental, and vision
The company pays the full premium for the employee and shares the cost for family members.
– Janus Soft pays 100% of the premium for the employee, and a generous contribution for
other family members, for health, dental, and vision insurance.
– Coverage is effective the first of the next month from the date of hire.
Coverage | Employee contribution | Employer contribution | Effective date
Employee only | 0% | 100% | 1st of the month after the date of hire
Spouse and children | 50% | 50% | 1st of the month after the date of hire
For candidates · 2026–27 Page 1
{HEADER}

02 · Time
Paid time off
Time away is part of the offer, not a balance you have to earn week by week.
– Unlimited paid time off, subject to meeting the minimum work-hours requirement.

03 · Future
Retirement
The 401(k) is a safe-harbor plan, and the match follows the employee's own contribution.
– The company matches 10% of the employee contribution.

06 · Protection
Disability and life insurance
Income protection and life insurance are company-paid on the schedule below.
– Janus Soft offers short-term disability, long-term disability, and life insurance.
Coverage | Employee contribution | Employer contribution | Effective date
Life insurance | 0% | 100% | 1st of the month, 60 days after the date of hire

08 · Tools
Health and technology
A yearly allowance covers the tools and programs people actually use.
– Janus Soft reimburses up to $1,000.00 per calendar year for laptops, gym memberships, cell phones, and tablets.
These terms are the 2026–27 benefits summary of Janus Soft Inc.
"""


def _client(text: str = "Short answer.") -> tuple[TestClient, RecordingLLM]:
    reset_rate_limit()
    reset_check_limit()
    app = create_app()
    client = TestClient(app)
    client.__enter__()
    llm = RecordingLLM(text=text)
    app.state.llm = llm
    app.state.embedder = RecordingEmbedder()
    return client, llm


def _login(client: TestClient) -> None:
    assert client.post("/admin/login", json={"email": "admin@janus-soft.com", "password": "correct-horse"}).status_code == 200


def _upload(client: TestClient, text: str = GUIDE) -> None:
    created = client.post("/admin/knowledge", data={"title": "Company Benefits", "audience": "public", "text": text})
    assert created.status_code == 200, created.text


def _ask(client: TestClient, message: str, **extra) -> dict:
    response = client.post("/public/chat", json={"message": message, "prior_codes": [], **extra})
    assert response.status_code == 200, response.text
    return response.json()


def test_a_guide_splits_into_titled_sections_without_page_chrome():
    sections = chunk_sections(GUIDE)
    headings = [section.heading for section in sections]
    assert headings == [
        None,
        "Medical, dental, and vision",
        "Paid time off",
        "Retirement",
        "Disability and life insurance",
        "Health and technology",
        None,
    ]
    care = sections[1]
    assert care.kicker == "01 · Care"
    assert "Page 1" not in care.body and "Prosperity" not in care.body
    # The wrapped PDF line is joined back into its bullet.
    assert any(line.endswith("for health, dental, and vision insurance.") for line in care.lines)
    # The closing note belongs to the document, not to the last section.
    assert "These terms" not in sections[5].body
    # A phrase repeated in the glance box and as a heading is content, not a page header.
    assert "Paid time off" in clean_lines(GUIDE)


def test_tables_read_as_sentences_and_blocks():
    lines = chunk_sections(GUIDE)[1].lines
    shown = blocks(lines)
    table = next(block for block in shown if block["kind"] == "table")
    assert table["header"][0] == "Coverage"
    assert table["rows"][1][:3] == ["Spouse and children", "50%", "50%"]
    assert "Spouse and children: employee contribution 50%, employer contribution 50%" in " ".join(readable(lines))


def test_docx_tables_keep_their_rows():
    document = Document()
    document.add_paragraph("Medical, dental, and vision coverage for every employee of the company and their families.")
    table = document.add_table(rows=2, cols=3)
    for row, values in zip(table.rows, (["Coverage", "Employee contribution", "Employer contribution"], ["Employee only", "0%", "100%"])):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("Coverage is effective the first of the next month from the date of hire, for all full-time staff.")
    buffer = BytesIO()
    document.save(buffer)
    text = extract_text("guide.docx", buffer.getvalue())
    assert "Coverage | Employee contribution | Employer contribution" in text
    assert "Employee only | 0% | 100%" in text
    assert text.index("Medical") < text.index("Employee only") < text.index("effective the first")


def test_health_benefits_returns_the_medical_section(db):
    client, llm = _client()
    _login(client)
    _upload(client)
    body = _ask(client, "show me the health benefits", explain=False)
    sections = [source["section"] for source in body["sources"]]
    assert sections[0] == "Medical, dental, and vision"
    assert body["knowledge_pending"] is True
    assert "Employee only: employee contribution 0%, employer contribution 100%" in body["answer"]
    assert "Spouse and children" in body["answer"]
    table = next(block for block in body["sources"][0]["blocks"] if block["kind"] == "table")
    assert table["rows"][0][0] == "Employee only"
    assert llm.prompts == []


def test_specific_questions_pick_their_own_section(db):
    client, _llm = _client()
    _login(client)
    _upload(client)
    cases = {
        "What is the 401k match?": "Retirement",
        "Is there life insurance?": "Disability and life insurance",
        "Do you reimburse gym memberships?": "Health and technology",
        "How much PTO do I get?": "Paid time off",
        "Is my family covered and what do I pay?": "Medical, dental, and vision",
    }
    for question, section in cases.items():
        body = _ask(client, question, explain=False)
        assert body["sources"], question
        assert body["sources"][0]["section"] == section, (question, [s["section"] for s in body["sources"]])
    missing = _ask(client, "What's the dress code?", explain=False)
    assert missing["sources"] == [] and missing["answer"] == "That is not in the knowledge base on file."


def test_a_general_benefits_question_lists_every_section(db):
    client, _llm = _client()
    _login(client)
    _upload(client)
    body = _ask(client, "Tell me about benefits", explain=False)
    assert body["overview"] is True
    assert [source["section"] for source in body["sources"]] == [
        "Medical, dental, and vision",
        "Paid time off",
        "Retirement",
        "Disability and life insurance",
        "Health and technology",
    ]
    assert body["answer"].startswith("Company Benefits covers 5 areas.")


def test_follow_ups_keep_the_topic_and_new_topics_search_jobs(db):
    client, _llm = _client()
    _login(client)
    _upload(client)
    history = [
        {"role": "user", "content": "Is there disability insurance?"},
        {"role": "assistant", "content": "Yes."},
    ]
    start = _ask(client, "and when does it start?", history=history, explain=False)
    assert start["sources"][0]["section"] == "Disability and life insurance"
    db.add(Job(requisition_code="A1007", title="Python Engineer", location="Chantilly, VA", status="open",
               description_text="Build services in Python.", must_have_skills=["Python"], nice_to_have_skills=[]))
    db.commit()
    python = _ask(client, "what about Python?", history=history, explain=False)
    assert python.get("sources") in (None, [])


def test_the_written_answer_is_dropped_when_it_adds_a_figure(db):
    client, llm = _client(text="The company pays 80% of family premiums.\n• Starts on day 1.")
    _login(client)
    _upload(client)
    reply = client.post("/public/knowledge/answer", json={"message": "Is my family covered?"})
    assert reply.status_code == 200
    answer = reply.json()["answer"]
    assert "80%" not in answer
    assert "Spouse and children: employee contribution 50%" in answer
    assert llm.prompts and "Medical, dental, and vision" in llm.prompts[0]["user"]

    llm.text = "**Yes.** Family members are covered.\n- Spouse and children pay 50% and the company pays 50%."
    good = client.post("/public/knowledge/answer", json={"message": "Is my family covered?"}).json()["answer"]
    assert good.startswith("Yes. Family members are covered.")
    assert "• Spouse and children pay 50%" in good
    assert "**" not in good


def test_written_answers_drop_not_specified_lines(db):
    client, llm = _client(
        text="Disability and life insurance are company paid.\n"
        "Life insurance: employee 0%, employer 100%.\n"
        "Long-term disability is offered but contribution details are not provided in the text.\n"
        "The provided sections do not list specific AWS job roles."
    )
    _login(client)
    _upload(client)
    answer = client.post("/public/knowledge/answer", json={"message": "Is there disability insurance?"}).json()["answer"]
    assert "Life insurance: employee 0%, employer 100%." in answer
    assert "not provided" not in answer
    assert "do not list" not in answer


def test_figures_match_reads_money_and_percent():
    passages = [{"excerpt": "Up to $1,000.00 per year. 50% employee.", "section": "Tools", "kicker": "08 · Tools"}]
    assert figures_match("Up to $1,000 a year, and you pay 50%.", passages, "gym?")
    assert not figures_match("Up to $2,000 a year.", passages, "gym?")


def test_a_role_and_a_benefit_in_one_question(db):
    client, _llm = _client()
    _login(client)
    _upload(client)
    db.add(Job(requisition_code="G2001", title="AWS Cloud Engineer", location="McLean, VA", status="open",
               description_text="Operate AWS accounts.", must_have_skills=["AWS"], nice_to_have_skills=[]))
    db.commit()
    question = "Does the AWS role include dental insurance?"
    assert wants_knowledge(question) and names_job(question)
    assert job_part(question) == "AWS role"
    assert job_part("python developer roles with good health insurance") == "python developer roles"
    assert job_part("any AWS jobs, and do they include dental?") == "AWS jobs"
    body = _ask(client, question, explain=False)
    assert body["sources"][0]["section"] == "Medical, dental, and vision"
    assert [job["requisition_code"] for job in body["jobs"]] == ["G2001"]


def test_soft_benefit_words_do_not_capture_job_searches():
    assert wants_knowledge("show me the health benefits")
    assert wants_knowledge("gym reimbursement")
    assert not wants_knowledge("health care data analyst jobs")
    assert not wants_knowledge("Java developer salary")


def test_quoted_overview_names_each_section():
    passages = [
        {"title": "Guide", "section": "Retirement", "excerpt": "The 401(k) is safe harbor.\n– 10% match.", "lines": ["The 401(k) is safe harbor.", "– 10% match."], "overview": True},
        {"title": "Guide", "section": "Paid time off", "excerpt": "Time away.\n– Unlimited.", "lines": ["Time away.", "– Unlimited."], "overview": True},
    ]
    text = quoted_answer(passages, overview=True)
    assert "• Retirement: The 401(k) is safe harbor." in text and "• Paid time off: Time away." in text


RESUME = (
    "Jordan Lee\nFront-end engineer in Chantilly, VA.\n"
    "Skills: React, Next.js, JavaScript, TypeScript, HTML, CSS, Figma.\n"
    "Built React and Next.js applications in TypeScript for six years, translating research into UI components "
    "and wireframes. Partnered with agile development teams to ship design systems.\n"
)


def _ui_job(db) -> Job:
    job = Job(
        requisition_code="A1001",
        title="UI/UX Developer",
        location="Chantilly, VA",
        status="open",
        description_text="UI work.",
        must_have_skills=["JavaScript"],
        nice_to_have_skills=[],
        skill_quotes={
            "must": [
                "Front-End Development: Hands-on experience with modern front-end technologies (React, Next.js) and core web languages (JavaScript, TypeScript, HTML, CSS).",
                "Cloud: Deploys services on AWS and Kubernetes.",
            ],
            "nice": ["Advanced prototyping expertise using industry-standard tools (Figma, Sketch, Adobe XD)."],
            "clearance": "Clearance: Active TS/SCI with Full Scope Polygraph.",
        },
    )
    db.add(job)
    db.commit()
    return job


def test_resume_check_lists_met_and_missing_requirements_without_storing(db):
    _ui_job(db)
    client, llm = _client()
    response = client.post(
        "/public/apply/check",
        data={"requisition_code": "A1001"},
        files={"file": ("jordan.txt", RESUME.encode(), "text/plain")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    items = body["required"]["items"]
    assert [item["label"] for item in items] == ["Front-End Development", "Cloud"]
    assert items[0]["met"] is True and items[1]["met"] is False
    assert {"AWS", "Kubernetes"} <= set(body["missing_tools"])
    assert body["required"]["hit"] == 1 and body["required"]["total"] == 2
    assert body["verdict"] == "partial"
    assert body["desired"]["total"] == 1
    tools = {tool["name"]: tool["have"] for tool in items[0]["tools"]}
    assert tools["React"] is True and tools["JavaScript"] is True
    assert llm.prompts == []
    assert db.query(Job).count() == 1
    from app.models import Candidate, JobApplication

    assert db.query(Candidate).count() == 0 and db.query(JobApplication).count() == 0


def test_resume_check_refuses_closed_jobs_and_bad_files(db):
    _ui_job(db)
    client, _llm = _client()
    closed = client.post("/public/apply/check", data={"requisition_code": "Z9"}, files={"file": ("a.txt", RESUME.encode(), "text/plain")})
    assert closed.status_code == 404
    empty = client.post("/public/apply/check", data={"requisition_code": "A1001"}, files={"file": ("a.txt", b"", "text/plain")})
    assert empty.status_code == 400
    for _ in range(20):
        client.post("/public/apply/check", data={"requisition_code": "A1001"}, files={"file": ("a.txt", RESUME.encode(), "text/plain")})
    limited = client.post("/public/apply/check", data={"requisition_code": "A1001"}, files={"file": ("a.txt", RESUME.encode(), "text/plain")})
    assert limited.status_code == 429
